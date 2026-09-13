import os
from datetime import datetime

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from database import create_tables, get_admin_chat_ids, get_weekly_user_report_data
from notifications import send_telegram_document

def build_docx_report(output_path: str) -> dict:
    """Haftalık raporu bir .docx dosyası olarak oluşturur."""
    data = get_weekly_user_report_data()
    simdi = datetime.now()

    doc = Document()

    # --- Başlık ---
    baslik = doc.add_heading("BIST Pivot Projeksiyon Sistemi", level=0)
    baslik.alignment = WD_ALIGN_PARAGRAPH.CENTER

    alt_baslik = doc.add_paragraph("Haftalık Kullanıcı Raporu")
    alt_baslik.alignment = WD_ALIGN_PARAGRAPH.CENTER
    alt_baslik.runs[0].bold = True
    alt_baslik.runs[0].font.size = Pt(14)

    tarih_p = doc.add_paragraph(simdi.strftime("%d.%m.%Y %H:%M"))
    tarih_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    # --- Özet İstatistikler ---
    doc.add_heading("Genel Durum", level=1)

    ozet_tablo = doc.add_table(rows=4, cols=2)
    ozet_tablo.style = "Light Grid Accent 1"

    ozet_veriler = [
        ("Toplam Kullanıcı", str(data["toplam_kullanici"])),
        ("Telegram Bağlı kullanıcı", str(data["telegram_bagli"])),
        ("Toplam Alert", str(data["toplam_alert"])),
        ("Aktif Alert", str(data["aktif_alert"])),
    ]
    for i, (etiket, deger) in enumerate(ozet_veriler):
        ozet_tablo.rows[i].cells[0].text = etiket
        ozet_tablo.rows[i].cells[1].text = deger

    doc.add_paragraph()

    # --- Son 7 Günde Kaydolan Kullanıcılar ---
    doc.add_heading("Son 7 Günde Kaydolan Kullanıcılar", level=1)

    yeniler = data["yeni_kullanicilar"]
    if yeniler:
        yeni_tablo = doc.add_table(rows=1, cols=3)
        yeni_tablo.style = "Light Grid Accent 1"
        basliklar = yeni_tablo.rows[0].cells
        basliklar[0].text = "Kullanıcı Adı"
        basliklar[1].text = "Ad-Soyad"
        basliklar[2].text = "Kayıt Zamanı"

        for username, first_name, last_name, created_at in yeniler:
            satir = yeni_tablo.add_row().cells
            satir[0].text = username
            satir[1].text = f"{first_name} {last_name}"
            satir[2].text = str(created_at)
    else:
        doc.add_paragraph("Bu hafta yeni kayıt bulunmamaktadır.")

    doc.add_paragraph()

    # --- Tüm Kullanıcılar ---
    doc.add_heading("Tüm Kullanıcılar", level=1)

    tum_tablo = doc.add_table(rows=1, cols=5)
    tum_tablo.style = "Light Grid Accent 1"
    basliklar = tum_tablo.rows[0].cells
    basliklar[0].text = "Kullanıcı Adı"
    basliklar[1].text = "Ad-Soyad"
    basliklar[2].text = "E-posta"
    basliklar[3].text = "Telegram"
    basliklar[4].text = "Alert"

    for username, first_name, last_name, email, tg, alert_sayisi in data["kullanici_ozeti"]:
        satir = tum_tablo.add_row().cells
        satir[0].text = username
        satir[1].text = f"{first_name} {last_name}"
        satir[2].text = email
        satir[3].text = "Bağlı" if tg else "Bağlı Değil"
        satir[4].text = str(alert_sayisi)

    # --- Dipnot ---
    doc.add_paragraph()
    dipnot = doc.add_paragraph(
        "Bu rapor otomatik olarak oluşturulmuştur | BIST Pivot Point Tabanlı Hisse Projeksiyon Sistemi ©2026"
    )
    dipnot.runs[0].font.size = Pt(9)
    dipnot.runs[0].font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    doc.save(output_path)
    return data

def send_weekly_report():
    """Word raporunu oluşturup admin'lere Telegram üzerinden gönderir."""
    create_tables()  # scripti self-contained yapar

    simdi = datetime.now()
    dosya_adi = f"haftalik_rapor_{simdi.strftime('%Y%m%d')}.docx"

    print("Rapor oluşturuluyor...")
    data = build_docx_report(dosya_adi)
    print(f"Rapor Oluşturuldu: {dosya_adi}")
    print(f"  Toplam Kullanıcı: {data['toplam_kullanici']}")
    print(f"  Son 7 günde Kaydolan Yeni Kullanıcı Sayısı: {len(data['yeni_kullanicilar'])}")

    admin_ids = get_admin_chat_ids()
    if not admin_ids:
        print("\nUYARI: Telegrama bağlı admin bulunamadı, rapor gönderilemedi.")
        return

    aciklama = (
        f"📊Haftalık Kullanıcı Raporu — {simdi.strftime('%d.%m.%Y')}\n"
        f"Toplam Kullanıcı: {data['toplam_kullanici']} | "
        f"Bu Hafta Yeni: {len(data['yeni_kullanicilar'])}"
    )

    for chat_id in admin_ids:
        if send_telegram_document(chat_id, dosya_adi, caption=aciklama):
            print(f"Rapor gönderildi (chat_id: {chat_id})")
        else:
            print(f"UYARI: Rapor gönderilemedi (chat_id: {chat_id})")

    # Gecici dosyayı temizle
    try:
        os.remove(dosya_adi)
    except OSError:
        pass


if __name__ == "__main__":
    send_weekly_report()